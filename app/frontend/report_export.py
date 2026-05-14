"""Word export and helpers (``python-docx`` is imported only when building .docx)."""

from __future__ import annotations

import json
import re
from io import BytesIO
from typing import Any

from app.frontend.meeting_registry import meeting_id_from_report_filename
from app.frontend.report_artifacts import job_paths_for_report, load_polished_markdown, merged_risk_view


def _add_mdish_block(doc: Any, text: str) -> None:
    """Rough Markdown to docx: headings and paragraphs (tables stay as plain lines)."""
    for raw in text.split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
        elif re.match(r"^[-*]\s+", line):
            doc.add_paragraph(re.sub(r"^[-*]\s+", "", line), style="List Bullet")
        else:
            doc.add_paragraph(line[:4000])


def build_meeting_docx_bytes(data: dict[str, Any], report_filename: str) -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt
    except ModuleNotFoundError as exc:
        raise ModuleNotFoundError(
            "Install Word export support: pip install python-docx"
        ) from exc

    paths = job_paths_for_report(report_filename)
    stem = meeting_id_from_report_filename(report_filename)
    md = load_polished_markdown(paths) or ""
    risk = merged_risk_view(data, paths)

    doc = Document()
    sty = doc.styles["Normal"]
    sty.font.name = "Calibri"
    sty.font.size = Pt(11)

    doc.add_heading(f"Meeting report - {stem}", 0)

    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph(str(data.get("summary") or "").strip()[:8000] or "(No summary in packaged JSON.)")

    doc.add_heading("Risk review", level=1)
    doc.add_paragraph("Needs human review: " + ("yes" if risk.get("needs_review") else "no"))
    risks = risk.get("risks") or []
    if isinstance(risks, list) and risks:
        for r in risks:
            doc.add_paragraph(str(r), style="List Bullet")
    else:
        doc.add_paragraph("(No risk bullets in standalone risk report or final JSON.)")

    doc.add_heading("Polished meeting report", level=1)
    if md.strip():
        _add_mdish_block(doc, md[:120_000])
    else:
        doc.add_paragraph(
            f"Expected file not found: `{paths.polished_md.as_posix()}`. Re-run the pipeline to regenerate it."
        )

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def risk_json_bytes(data: dict[str, Any], report_filename: str) -> bytes:
    paths = job_paths_for_report(report_filename)
    payload = merged_risk_view(data, paths)
    return json.dumps(payload, indent=2, ensure_ascii=False).encode("utf-8")
